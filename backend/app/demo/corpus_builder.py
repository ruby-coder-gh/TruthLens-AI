"""Builds the demo corpus's binary artifacts (PDF/DOCX) from tracked sources.

The corpus is authored as plain markdown/CSV (`app/demo/corpus/*`, tracked in
git) and rendered into real PDF/DOCX files at seed time — see `app.demo.seed`.
Rendering rather than committing binaries keeps the fictional "Northwind
Renewables" corpus reviewable as text and its planted contradictions easy to
audit in a diff.
"""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
import markdown
from docx import Document as DocxDocument

CORPUS_DIR = Path(__file__).parent / "corpus"
_NUMBERED_ITEM = re.compile(r"\d+\. ")

MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "md": "text/markdown",
    "csv": "text/csv",
}


def load_manifest() -> dict[str, Any]:
    """Load the demo corpus manifest (doc list, suggested questions, contradictions)."""
    return json.loads((CORPUS_DIR / "manifest.json").read_text(encoding="utf-8"))


def build_pdf(markdown_path: Path, dest_path: Path) -> None:
    """Render a markdown source file to a paginated PDF at *dest_path*.

    Uses PyMuPDF's `Story` layout engine over the markdown-rendered HTML, so
    headings and the annual report's table survive as real PDF content
    (not a screenshot), which the Contradiction Radar and source viewer both
    need to work against.
    """
    html = markdown.markdown(markdown_path.read_text(encoding="utf-8"), extensions=["tables"])
    story = fitz.Story(html=html)
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (36, 36, -36, -36)

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    writer = fitz.DocumentWriter(str(dest_path))
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()


def build_docx(markdown_path: Path, dest_path: Path) -> None:
    """Render a markdown source file to a DOCX at *dest_path*.

    Naive line-based mapping (headings / bullets / numbered items / paragraphs,
    `**bold**` as bold runs) — good enough fidelity for a demo corpus; not a
    general markdown-to-DOCX converter. Hard-wrapped source lines rejoin into
    one DOCX paragraph; a line opening with `**` (a memo's "**To:** ..." header)
    starts its own.
    """
    blocks: list[tuple[str | None, list[str]]] = []  # (paragraph style, source lines)
    open_block = False
    for raw_line in markdown_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        numbered = _NUMBERED_ITEM.match(line)
        if not line:
            open_block = False
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append((f"Heading {level}", [line[level:].strip()]))
            open_block = False
        elif line.startswith("- "):
            blocks.append(("List Bullet", [line[2:]]))
            open_block = True
        elif numbered:
            blocks.append(("List Number", [line[numbered.end():]]))
            open_block = True
        elif open_block and not line.startswith("**"):
            blocks[-1][1].append(line)
        else:
            blocks.append((None, [line]))
            open_block = True

    doc = DocxDocument()
    for style, lines in blocks:
        paragraph = doc.add_paragraph(style=style)
        for k, part in enumerate(" ".join(lines).split("**")):
            if part:
                run = paragraph.add_run(part)
                if k % 2:
                    run.bold = True

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest_path))


def build_corpus_files(dest_dir: Path) -> list[dict[str, Any]]:
    """Materialise every manifest doc as a real file under *dest_dir*.

    Returns the manifest's `docs` list, each entry augmented with the local
    `path` (Path) it was written to and its `mime_type`.
    """
    manifest = load_manifest()
    dest_dir.mkdir(parents=True, exist_ok=True)

    built: list[dict[str, Any]] = []
    for entry in manifest["docs"]:
        src = CORPUS_DIR / entry["source"]
        stem = src.stem
        fmt = entry["format"]

        if fmt == "pdf":
            dest = dest_dir / f"{stem}.pdf"
            build_pdf(src, dest)
        elif fmt == "docx":
            dest = dest_dir / f"{stem}.docx"
            build_docx(src, dest)
        elif fmt in ("md", "csv"):
            dest = dest_dir / src.name
            shutil.copyfile(src, dest)
        else:
            raise ValueError(f"Unknown demo corpus format: {fmt!r}")

        built.append({**entry, "path": dest, "mime_type": MIME_TYPES[fmt]})

    return built
