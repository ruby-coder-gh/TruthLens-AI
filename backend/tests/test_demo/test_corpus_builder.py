"""Tests for the demo corpus builder (PDF/DOCX generation from tracked sources)."""

from __future__ import annotations

import unicodedata
from pathlib import Path

import fitz
from docx import Document as DocxDocument

from app.demo.corpus_builder import (
    CORPUS_DIR,
    build_corpus_files,
    build_docx,
    build_pdf,
    load_manifest,
)


def test_load_manifest_has_expected_shape():
    manifest = load_manifest()
    assert manifest["workspace_name"]
    assert len(manifest["docs"]) == 6
    assert len(manifest["suggested_questions"]) == 6
    assert len(manifest["planted_contradictions"]) >= 4
    for doc in manifest["docs"]:
        assert doc["format"] in {"pdf", "docx", "md", "csv"}
        assert (CORPUS_DIR / doc["source"]).exists()


def test_build_pdf_produces_a_readable_multi_section_pdf(tmp_path: Path):
    dest = tmp_path / "annual-report-2025.pdf"
    build_pdf(CORPUS_DIR / "annual-report-2025.md", dest)

    assert dest.exists()
    pdf = fitz.open(dest)
    try:
        assert pdf.page_count >= 1
        # NFKC undoes PyMuPDF's own ligature substitution (e.g. "fi" -> a
        # single glyph) — the actual ingestion path normalizes the same way
        # (app.ingestion.loader._load_pdf), covered by its own regression test.
        full_text = unicodedata.normalize(
            "NFKC", "\n".join(page.get_text() for page in pdf)
        )
        assert "Dana Whitfield" in full_text
        assert "412" in full_text
    finally:
        pdf.close()


def test_build_docx_produces_a_readable_document(tmp_path: Path):
    dest = tmp_path / "board-memo-aurora.docx"
    build_docx(CORPUS_DIR / "board-memo-aurora.md", dest)

    assert dest.exists()
    doc = DocxDocument(str(dest))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Aurora" in full_text
    assert "Q1 2028" in full_text or "first quarter of 2028" in full_text.lower() or "2028" in full_text


def test_build_docx_renders_markdown_paragraphs_and_bold(tmp_path: Path):
    """Real demo seed: the radar quoted "…estimated at **€1.1 billion**…" verbatim, and each
    hard-wrapped source line had become its own DOCX paragraph."""
    dest = tmp_path / "board-memo-aurora.docx"
    build_docx(CORPUS_DIR / "board-memo-aurora.md", dest)

    paragraphs = DocxDocument(str(dest)).paragraphs
    texts = [p.text for p in paragraphs]
    assert not [t for t in texts if "**" in t]
    budget = next(p for p in paragraphs if p.text.startswith("Total project capital expenditure"))
    assert budget.text.endswith("though this will be kept under review.")
    assert [r.text for r in budget.runs if r.bold] == ["€1.1 billion"]
    assert "To: Board of Directors" in texts
    risk = next(t for t in texts if t.startswith("Fabrication yard capacity."))
    assert risk.endswith("remaining monopiles not yet started.")


def test_build_corpus_files_writes_every_manifest_doc_with_expected_extension(tmp_path: Path):
    built = build_corpus_files(tmp_path)
    manifest = load_manifest()

    assert len(built) == len(manifest["docs"])
    for entry in built:
        assert entry["path"].exists()
        assert entry["path"].suffix.lstrip(".") == (
            "md" if entry["format"] == "md" else entry["format"]
        )
