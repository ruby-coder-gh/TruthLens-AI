"""Tests for document loader."""

from __future__ import annotations

from pathlib import Path

import pytest

from app.ingestion.loader import load


@pytest.mark.asyncio
async def test_load_txt(sample_txt_path: Path):
    """Test loading a text file."""
    pages = await load(sample_txt_path, "text/plain")
    assert len(pages) > 0
    assert "sample text" in pages[0]["text"]
    assert pages[0]["page_number"] is None


@pytest.mark.asyncio
async def test_load_nonexistent():
    """Test loading non-existent file raises error."""
    with pytest.raises(FileNotFoundError):
        await load(Path("nonexistent.txt"), "text/plain")


@pytest.mark.asyncio
async def test_load_unsupported_type():
    """Test loading unsupported mime type."""
    path = Path("test_data/test.xyz")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("test")
    with pytest.raises(ValueError, match="Unsupported mime type"):
        await load(path, "application/octet-stream")
    path.unlink()


@pytest.mark.asyncio
async def test_load_empty_txt():
    """Test loading empty text file."""
    path = Path("test_data/empty.txt")
    path.write_text("")
    pages = await load(path, "text/plain")
    assert len(pages) == 0
    path.unlink()


@pytest.mark.asyncio
async def test_load_csv():
    """Test loading a CSV file."""
    path = Path("test_data/test.csv")
    path.write_text("name,age\nAlice,30\nBob,25\n")
    pages = await load(path, "text/csv")
    assert len(pages) > 0
    assert "Alice" in pages[0]["text"]
    path.unlink()


@pytest.mark.asyncio
async def test_load_csv_one_chunk_per_row(tmp_path: Path):
    """C7: CSV rows become independently retrievable — one page per row, with
    the header repeated as 'Column: value; …' text (not one page for the whole file)."""
    path = tmp_path / "rows.csv"
    path.write_text("project,status\nAurora,construction\nKestrel Ridge,commissioned\n")
    pages = await load(path, "text/csv")

    prefix = "rows.csv. Table with columns: project, status.\n"
    assert len(pages) == 2
    assert pages[0]["text"] == prefix + "project: Aurora; status: construction"
    assert pages[1]["text"] == prefix + "project: Kestrel Ridge; status: commissioned"


@pytest.mark.asyncio
async def test_load_csv_row_chunks_are_prefixed_with_filename_and_description(tmp_path: Path):
    """BUG-7: a bare 'col: value; col: value' row scored near-zero against a
    natural-language question on the real reranker ("which projects are
    under construction" vs "name: Aurora; ...; status: construction" scored
    0.0004, threshold 0.35). Prefixing every row with the filename and a
    short table description gave every row real lexical/semantic overlap
    with a natural-language question about the table (0.39-0.89 in a repro
    against the real BAAI/bge-reranker-v2-m3 model, see FIX-round2 report)."""
    path = tmp_path / "projects.csv"
    path.write_text("name,status\nAurora,construction\n")
    pages = await load(path, "text/csv")

    assert pages[0]["text"].startswith("projects.csv. Table with columns: name, status.\n")


@pytest.mark.asyncio
async def test_load_csv_uses_its_own_leading_comment_line_as_the_description(tmp_path: Path):
    """BUG-7: a CSV can opt into a stronger, hand-written table description
    (real domain phrasing scored far higher than the generic column-name
    fallback in the real-reranker repro) via a `# ...` first line -- which
    is not treated as a data row."""
    path = tmp_path / "projects.csv"
    path.write_text(
        "# Renewable energy project pipeline and construction status.\n"
        "name,status\nAurora,construction\n"
    )
    pages = await load(path, "text/csv")

    assert len(pages) == 1
    assert pages[0]["text"] == (
        "projects.csv. Renewable energy project pipeline and construction status.\n"
        "name: Aurora; status: construction"
    )


@pytest.mark.asyncio
async def test_load_csv_adds_a_summary_chunk_per_low_cardinality_column(tmp_path: Path):
    """R3-4: aggregation questions ("which projects are under construction")
    only see the top reranked *row* chunks, silently missing rows past that
    cutoff. One extra chunk per low-cardinality column (<= 8 distinct values)
    lists every matching row's name together so the full group survives in
    a single chunk."""
    path = tmp_path / "pipeline.csv"
    path.write_text(
        "name,status\n"
        "Aurora,construction\n"
        "Fjellheim,construction\n"
        "Solheim,construction\n"
        "Ashford,construction\n"
        "Lindholm,construction\n"
        "Kestrel,permitting\n"
    )
    pages = await load(path, "text/csv")

    row_pages = [p for p in pages if p["metadata"].get("row_index") is not None]
    summary_pages = [p for p in pages if "summary_column" in p["metadata"]]
    assert len(row_pages) == 6

    construction = next(p for p in summary_pages if p["metadata"]["summary_value"] == "construction")
    assert construction["metadata"]["summary_column"] == "status"
    assert construction["text"].startswith("Pipeline. ")
    assert construction["text"].endswith("status = Construction: Aurora, Fjellheim, Solheim, Ashford, Lindholm")

    permitting = next(p for p in summary_pages if p["metadata"]["summary_value"] == "permitting")
    assert permitting["text"].endswith("status = Permitting: Kestrel")


@pytest.mark.asyncio
async def test_load_csv_skips_summary_for_high_cardinality_columns(tmp_path: Path):
    """A column where (almost) every row has a distinct value (e.g. a
    numeric measurement) doesn't get a summary chunk per value -- that would
    be one chunk per row all over again, not a useful rollup."""
    rows = "\n".join(f"Project{i},{i*10}" for i in range(10))
    path = tmp_path / "pipeline.csv"
    path.write_text(f"name,mw\n{rows}\n")
    pages = await load(path, "text/csv")

    assert not any("summary_column" in p["metadata"] for p in pages)


@pytest.mark.asyncio
async def test_load_csv_skips_summary_for_small_tables(tmp_path: Path):
    """A handful of rows all fit comfortably in top-k retrieval already --
    no rollup chunk needed (also keeps toy/unit-test CSVs summary-free)."""
    path = tmp_path / "pipeline.csv"
    path.write_text("name,status\nAurora,construction\nKestrel,commissioned\n")
    pages = await load(path, "text/csv")

    assert not any("summary_column" in p["metadata"] for p in pages)


@pytest.mark.asyncio
async def test_load_csv_summary_title_comes_from_the_filename(tmp_path: Path):
    path = tmp_path / "project-pipeline.csv"
    path.write_text("name,status\n" + "\n".join(f"P{i},construction" for i in range(5)))
    pages = await load(path, "text/csv")

    summary = next(p for p in pages if "summary_column" in p["metadata"])
    assert summary["text"].startswith("Project Pipeline. ")


@pytest.mark.asyncio
async def test_load_csv_titles_rows_and_summaries_with_the_original_filename(tmp_path: Path):
    """QA4: stored files are named by UUID, so the summary chunk read
    "Bea4A6B7 9718 …" and was never retrieved. The ingestion pipeline passes
    the user-facing filename through."""
    path = tmp_path / "bea4a6b7-9718-4c2e.csv"
    path.write_text("name,status\n" + "\n".join(f"P{i},construction" for i in range(5)))
    pages = await load(path, "text/csv", "Northwind Renewables — Project Pipeline.csv")

    summary = next(p for p in pages if "summary_column" in p["metadata"])
    assert summary["text"].startswith("Northwind Renewables — Project Pipeline. ")
    row = next(p for p in pages if p["metadata"].get("row_index") == 0)
    assert row["text"].startswith("Northwind Renewables — Project Pipeline.csv. ")
    assert "bea4a6b7" not in summary["text"].lower()


@pytest.mark.asyncio
async def test_load_real_demo_pipeline_csv_groups_every_construction_project():
    """The exact R3-4 repro against the real demo corpus file: 'which
    projects are under construction?' must be able to find all 5 (not 4)."""
    path = Path("app/demo/corpus/project-pipeline.csv")
    pages = await load(path, "text/csv")

    construction = next(
        p for p in pages
        if p["metadata"].get("summary_column") == "status" and p["metadata"]["summary_value"] == "construction"
    )
    names = construction["text"].rsplit("\n", 1)[-1].split(": ", 1)[1].split(", ")
    assert set(names) == {"Aurora", "Fjellheim Repowering", "Solheim Solar Park", "Ashford Solar", "Lindholm Solar"}


@pytest.mark.asyncio
async def test_load_json_object_flattens_to_key_path_value_lines(tmp_path: Path):
    """R2-5: a JSON object is flattened to 'key.path: value' lines instead of
    raising "Unsupported mime type" (the upload UI already advertises JSON)."""
    path = tmp_path / "facts.json"
    path.write_text('{"company": {"name": "Northwind", "founded": 2015}, "active": true}')
    pages = await load(path, "application/json")

    assert len(pages) == 1
    lines = pages[0]["text"].splitlines()
    assert "company.name: Northwind" in lines
    assert "company.founded: 2015" in lines
    assert "active: True" in lines
    assert pages[0]["page_number"] is None


@pytest.mark.asyncio
async def test_load_json_array_is_one_page_per_item(tmp_path: Path):
    """R2-5: a JSON array of records gets the same per-row retrievability as
    CSV -- one page per item, so a single record stays independently findable."""
    path = tmp_path / "rows.json"
    path.write_text(
        '[{"project": "Aurora", "status": "construction"},'
        ' {"project": "Kestrel Ridge", "status": "commissioned"}]'
    )
    pages = await load(path, "application/json")

    assert len(pages) == 2
    assert pages[0]["text"] == "project: Aurora\nstatus: construction"
    assert pages[1]["text"] == "project: Kestrel Ridge\nstatus: commissioned"


@pytest.mark.asyncio
async def test_load_json_invalid_raises_value_error(tmp_path: Path):
    """R2-5: malformed JSON fails ingestion with a clear error, not a bare
    "Unsupported mime type" or an unhandled exception."""
    path = tmp_path / "bad.json"
    path.write_text("{not valid json")
    with pytest.raises(ValueError, match="Invalid JSON"):
        await load(path, "application/json")


@pytest.mark.asyncio
async def test_load_json_empty_object_returns_no_pages(tmp_path: Path):
    path = tmp_path / "empty.json"
    path.write_text("{}")
    pages = await load(path, "application/json")
    assert pages == []


@pytest.mark.asyncio
async def test_load_pdf_reverses_ligature_glyphs():
    """PyMuPDF's `Story` HTML layout (used to build the demo corpus PDFs) can
    substitute a single ligature glyph for an "fi"/"fl" letter pair (e.g.
    "Whitfield" -> "Whitﬁeld", U+FB01). Loading must return the plain ASCII
    the document actually says, or BM25 tokenization and citation excerpts
    silently stop matching a name typed normally."""
    import fitz

    path = Path("test_data/ligature.pdf")
    path.parent.mkdir(parents=True, exist_ok=True)
    story = fitz.Story(html="<p>Dana Whitfield is the Chief Executive Officer.</p>")
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (36, 36, -36, -36)
    writer = fitz.DocumentWriter(str(path))
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()

    pages = await load(path, "application/pdf")

    assert len(pages) > 0
    assert "Whitfield" in pages[0]["text"]
    assert "ﬁ" not in pages[0]["text"]
    path.unlink()


@pytest.mark.asyncio
async def test_load_pdf_returns_one_clean_paragraph_per_text_block(tmp_path: Path):
    """Plain `page.get_text()` hard-breaks every rendered line, so a heading, a
    table row and the next paragraph's prose ran together into one pseudo-
    sentence ("Sincerely, Dana Whitfield Chief Executive Officer Financial
    Highlights Metric 2025 ...") that the Contradiction Radar sent to NLI."""
    import fitz

    prose = (
        "Revenue in 2025 was €412 million, up from €356 million in 2024, driven by a full year "
        "of output from the Kestrel Ridge onshore expansion and higher merchant power prices."
    )
    html = (
        "<h2>Financial Highlights</h2>"
        f"<p>{prose}</p>"
        "<table><tr><td>Revenue</td><td>€412M</td></tr></table>"
    )
    path = tmp_path / "blocks.pdf"
    story = fitz.Story(html=html)
    writer = fitz.DocumentWriter(str(path))
    more = True
    while more:
        device = writer.begin_page(fitz.paper_rect("a5"))
        more, _ = story.place(fitz.paper_rect("a5") + (36, 36, -36, -36))
        story.draw(device)
        writer.end_page()
    writer.close()

    pages = await load(path, "application/pdf")

    assert pages[0]["page_number"] == 1
    assert pages[0]["text"].split("\n\n") == ["Financial Highlights", prose, "Revenue €412M"]


@pytest.mark.asyncio
async def test_load_docx_separates_paragraphs_with_blank_lines(tmp_path: Path):
    """Same contract as PDF blocks: a blank line is a hard paragraph break, so a heading is
    never read as the start of the sentence below it (a single newline may be a soft wrap)."""
    from docx import Document as DocxDocument

    path = tmp_path / "memo.docx"
    doc = DocxDocument()
    doc.add_heading("Outlook", level=2)
    doc.add_paragraph("2026 capital expenditure guidance is €640 million.")
    doc.save(str(path))

    pages = await load(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    assert pages[0]["text"] == "Outlook\n\n2026 capital expenditure guidance is €640 million."


@pytest.mark.asyncio
async def test_load_markdown_separates_blocks_with_blank_lines(tmp_path: Path):
    path = tmp_path / "update.md"
    path.write_text("## Project update\nAurora is expected to commission in\nthe third quarter of 2027.\n")

    pages = await load(path, "text/markdown")

    assert pages[0]["text"] == "Project update\n\nAurora is expected to commission in\nthe third quarter of 2027."


@pytest.mark.asyncio
async def test_load_pdf_rejoins_words_hyphenated_across_lines(tmp_path: Path):
    import fitz

    path = tmp_path / "hyphen.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 100), "Turbine foundation fabri-\ncation began in the fourth quarter.")
    doc.save(path)
    doc.close()

    pages = await load(path, "application/pdf")

    assert pages[0]["text"] == "Turbine foundation fabrication began in the fourth quarter."


def test_block_paragraph_keeps_hyphen_of_compound_split_before_capital():
    from app.ingestion.loader import _block_paragraph

    assert _block_paragraph("Reports Fourth-\nQuarter and Full-Year") == "Reports Fourth-Quarter and Full-Year"
    assert _block_paragraph("fabri-\ncation of Q4-\n2025 results") == "fabrication of Q4-2025 results"
